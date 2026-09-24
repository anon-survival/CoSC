import torch
import util
import torch.nn.functional as F
import copy
from openpyxl import load_workbook
import time
import optim
import torch.nn as nn
import os
import random
import numpy as np

DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")

def safe_logit(x, eps=1e-6):
    return torch.log((x + eps) / (1 - x + eps))

def normalized_rmst(cdf_matrix, time_grid, tau=None):
    """Compute RMST/tau on a shared, sorted time grid."""
    if cdf_matrix.ndim != 2 or time_grid.ndim != 1:
        raise ValueError("cdf_matrix must be 2-D and time_grid must be 1-D")
    if cdf_matrix.shape[1] != time_grid.numel():
        raise ValueError("cdf_matrix columns must match time_grid")
    if torch.any(time_grid[1:] < time_grid[:-1]):
        raise ValueError("time_grid must be sorted")

    if tau is None:
        tau = time_grid[-1]
    tau = torch.as_tensor(tau, dtype=time_grid.dtype, device=time_grid.device)
    if tau <= 0:
        raise ValueError("tau must be positive")

    survival = 1 - cdf_matrix
    before_tau = time_grid < tau
    times = time_grid[before_tau]
    survival_before = survival[:, before_tau]

    right = torch.searchsorted(time_grid, tau).clamp(max=time_grid.numel() - 1)
    left = (right - 1).clamp(min=0)
    left_t, right_t = time_grid[left], time_grid[right]
    fraction = ((tau - left_t) / (right_t - left_t).clamp_min(torch.finfo(time_grid.dtype).eps)).clamp(0, 1)
    survival_tau = survival[:, left] + fraction * (survival[:, right] - survival[:, left])
    survival_tau = torch.where(tau > time_grid[-1], survival[:, -1], survival_tau)

    integration_times = torch.cat([time_grid.new_zeros(1), times, tau.reshape(1)])
    integration_survival = torch.cat([survival.new_ones((survival.shape[0], 1)), survival_before, survival_tau.unsqueeze(1)], dim=1)

    return torch.trapezoid(integration_survival, integration_times, dim=1) / tau

def postprocessing(args, cdf, is_dead, device='cpu', max_iters=10000, tol=1e-8, patience=100):
    EPS = 1e-8
    order = torch.argsort(cdf)
    cdf = cdf[order]
    cdf = cdf.unsqueeze(1)
    is_dead = is_dead[order].unsqueeze(1)
    N = cdf.shape[0]

    # Initialize learnable parameters
    a0_raw = torch.nn.Parameter(torch.tensor(0.0, device=device))
    b0 = torch.nn.Parameter(torch.tensor(0.0, device=device))
    alpha_raw = torch.nn.Parameter(torch.tensor(0.0, device=device))
    optimizer = torch.optim.Adam([a0_raw, b0, alpha_raw], lr=0.05)

    best_ks = float('inf')  # Track the best KS value
    best_params = None
    patience_counter = 0  # Initialize patience counter

    start_time = time.time()
    for iter in range(max_iters):

        with torch.set_grad_enabled(True):
            is_alive = (1 - is_dead).float()
            F_sorted = torch.sigmoid(torch.exp(a0_raw) * safe_logit(cdf) + b0) ** torch.exp(alpha_raw)

            denom = 1 - F_sorted + EPS
            weight = is_alive / denom
            F_weight = F_sorted * weight

            cum_weight = torch.cumsum(weight, dim=0)
            cum_F_weight = torch.cumsum(F_weight, dim=0)

            cum_weight_shifted = F.pad(cum_weight[:-1], (0, 0, 1, 0), value=0.0)
            cum_F_weight_shifted = F.pad(cum_F_weight[:-1], (0, 0, 1, 0), value=0.0)

            ecdf_cens = F_sorted * cum_weight_shifted - cum_F_weight_shifted
            ecdf_cens = torch.clamp(ecdf_cens, 0, N)

            ecdf_dead = torch.cumsum(is_dead, dim=0)
            ecdf_upper = (ecdf_dead + ecdf_cens) / N
            ecdf_upper = torch.clamp(ecdf_upper, 0, 1)
            ecdf_lower = ecdf_upper - is_dead / N

            KS_upper = torch.abs(ecdf_upper - F_sorted)
            KS_lower = torch.abs(ecdf_lower - F_sorted)
            KS_error = torch.max(torch.concat([KS_upper, KS_lower], dim=1), dim=1).values
            KS = torch.max(KS_error)

            print(f"KS: {KS.item():.6f}, Iteration: {iter+1}/{max_iters}", end="\r")

            optimizer.zero_grad()
            KS.backward()
            optimizer.step()

            # Gradient tolerance check
            grad_norm = torch.norm(torch.cat([a0_raw.grad.view(1), b0.grad.view(1), alpha_raw.grad.view(1)]))

            if grad_norm < tol:
                print(f"\nGradient norm below tolerance: {grad_norm:.6f}. Stopping early at iteration {iter+1}.")
                break

            # Early stopping check
            if torch.isfinite(KS):
                if KS.item() < best_ks:
                    best_ks = KS.item()
                    best_params = (a0_raw.clone(), b0.clone(), alpha_raw.clone())
                    patience_counter = 0
                else:
                    patience_counter += 1
                    if patience_counter >= patience:
                        print(f"\nEarly stopping at iteration {iter+1}. Best KS: {best_ks:.6f}")
                        break
            else:
                print("\nNon-finite KS detected. Restoring best and stopping.")
                break

    end_time = time.time()
    print("KSP time:", end_time - start_time)
    # workbook = load_workbook(filename='./ksp_time.xlsx')
    # sheet = workbook.active
    # last_row = sheet.max_row
    # sheet.cell(row=last_row+1, column=1, value=(end_time-start_time))
    # sheet.cell(row=last_row+1, column=2, value=(f'KSP_{args.dataset}_{args.model_dist}'))
    # workbook.save('./ksp_time.xlsx')

    # workbook = load_workbook(filename='./total_iter.xlsx')
    # sheet = workbook.active
    # last_row = sheet.max_row
    # sheet.cell(row=last_row+1, column=1, value=iter+1)
    # sheet.cell(row=last_row+1, column=2, value=(f'KSP_{args.dataset}_{args.model_dist}'))
    # workbook.save('./total_iter.xlsx')

    # Restore best parameters
    a0_raw, b0, alpha_raw = best_params
    a0 = torch.exp(a0_raw).item()
    b0 = b0.item()
    alpha = torch.exp(alpha_raw).item()
    print("Final parameters after early stopping:")
    print("a0:", a0)
    print("b0:", b0)
    print("alpha:", alpha)

    return a0, b0, alpha

def ecdf(t, cdf, is_dead, weights=None, eps=1e-8):
    """
    t: 변환할 CDF 값, [N, T] 또는 임의 shape
    cdf: calibration PIT, [N_cal]
    is_dead: calibration event, [N_cal]
    weights: calibration weights, [N_cal] 또는 [N, N_cal]
    """
    original_shape = t.shape
    query = t.reshape(-1)
    observed = cdf.reshape(-1)
    event = is_dead.reshape(-1).float()
    if observed.numel() != event.numel():
        raise ValueError("cdf and is_dead must have the same number of elements")

    order = torch.argsort(observed)
    observed = observed[order]
    event = event[order]
    censor = 1 - event
    inverse_survival = censor / (1 - observed).clamp_min(eps)
    indices = torch.searchsorted(observed, query.contiguous(), right=True)

    def prefix_sum(values):
        zero_shape = (*values.shape[:-1], 1)
        return torch.cat([values.new_zeros(zero_shape), values.cumsum(dim=-1)], dim=-1)

    if weights is None or weights.ndim == 1:
        if weights is None:
            sorted_weights = torch.ones_like(observed)
        else:
            if weights.numel() != observed.numel():
                raise ValueError("1-D weights must match the calibration CDF length")
            sorted_weights = weights.reshape(-1)[order]

        event_prefix = prefix_sum(sorted_weights * event)
        censor_prefix = prefix_sum(sorted_weights * inverse_survival)
        censor_cdf_prefix = prefix_sum(sorted_weights * inverse_survival * observed)
        transformed = (
            event_prefix[indices]
            + query * censor_prefix[indices]
            - censor_cdf_prefix[indices]
        ) / sorted_weights.sum().clamp_min(eps)
    else:
        if t.ndim != 2 or weights.ndim != 2:
            raise ValueError("subject-specific weights require 2-D t and weights")
        if weights.shape != (t.shape[0], observed.numel()):
            raise ValueError("weights must have shape [n_subjects, n_calibration]")

        sorted_weights = weights[:, order]
        event_prefix = prefix_sum(sorted_weights * event)
        censor_prefix = prefix_sum(sorted_weights * inverse_survival)
        censor_cdf_prefix = prefix_sum(sorted_weights * inverse_survival * observed)
        row_indices = torch.arange(t.shape[0], device=t.device).repeat_interleave(t.shape[1])
        transformed = (
            event_prefix[row_indices, indices]
            + query * censor_prefix[row_indices, indices]
            - censor_cdf_prefix[row_indices, indices]
        ) / sorted_weights.sum(dim=1)[row_indices].clamp_min(eps)

    return transformed.clamp(0, 1).reshape(original_shape)

def marginal_censoring_km(observed, is_dead, query, left_limit=False, eps=1e-6):
    """Evaluate the marginal KM survival function of the censoring variable."""
    observed = observed.reshape(-1)
    censor_event = (1 - is_dead.reshape(-1).float())
    order = torch.argsort(observed)
    sorted_observed = observed[order]
    sorted_censor_event = censor_event[order]

    unique_times, inverse, counts = torch.unique_consecutive(
        sorted_observed, return_inverse=True, return_counts=True
    )
    censor_counts = sorted_observed.new_zeros(unique_times.numel())
    censor_counts.scatter_add_(0, inverse, sorted_censor_event)
    at_risk = observed.numel() - torch.cat(
        [counts.new_zeros(1), counts.cumsum(0)[:-1]]
    )
    survival_after = torch.cumprod(
        1 - censor_counts / at_risk.to(observed.dtype), dim=0
    )
    survival_prefix = torch.cat(
        [survival_after.new_ones(1), survival_after]
    )
    indices = torch.searchsorted(
        unique_times, query.contiguous(), right=not left_limit
    )
    return survival_prefix[indices].clamp_min(eps)
